# -*- coding: utf-8 -*-
"""Semanas 18-19: modelo de tres estados financieros con escenarios.

Todo esta enlazado por formulas vivas: cambia un supuesto y el modelo entero
reacciona, incluida la comprobacion de que el balance cuadra.
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
from openpyxl import Workbook
from openpyxl.styles import Alignment
from estilo_excel import (titulo, seccion, cabecera, fila_datos, nota, anchos, guardar,
                          F_NOTA, F_TOTAL, R_INPUT, R_TOTAL, BORDE,
                          NUM_MILES, NUM_DEC, NUM_PCT, NUM_PCT2, VERDE, ROJO)
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font

ANIOS = 5


def hoja_supuestos(wb):
    ws = wb.create_sheet('Supuestos')
    anchos(ws, 42, 14, 8)
    f = titulo(ws, 'SUPUESTOS DEL MODELO  ·  celdas amarillas = editables', 8)

    f = seccion(ws, f, 'Selector de escenario')
    f = fila_datos(ws, f, 'Escenario activo (1=Base  2=Optimista  3=Pesimista)',
                   [1], fmt='0', input_=True)
    escenario_celda = 'B%d' % (f - 1)
    ws['B%d' % (f - 1)].alignment = Alignment(horizontal='center')
    f = nota(ws, f, 'Cambia solo este numero: todo el modelo se recalcula.')
    f += 1

    f = seccion(ws, f, 'Tabla de escenarios')
    f = cabecera(ws, f, ['Supuesto', 'Base', 'Optimista', 'Pesimista'])
    inicio_tabla = f
    escenarios = [
        ('Crecimiento de ventas',      0.10, 0.18, 0.02, NUM_PCT),
        ('Margen EBITDA',              0.20, 0.24, 0.15, NUM_PCT),
        ('Tasa de impuestos',          0.25, 0.25, 0.25, NUM_PCT),
        ('Tasa de interes de la deuda', 0.05, 0.045, 0.07, NUM_PCT2),
    ]
    for etiqueta, base, opt, pes, fmt in escenarios:
        f = fila_datos(ws, f, etiqueta, [base, opt, pes], fmt=fmt, input_=True)
    fin_tabla = f - 1
    f += 1

    f = seccion(ws, f, 'Supuestos activos (calculados con CHOOSE / ELEGIR)')
    activos = []
    for i, (etiqueta, _, _, _, fmt) in enumerate(escenarios):
        fila_origen = inicio_tabla + i
        formula = '=IFERROR(CHOOSE($%s,B%d,C%d,D%d),"Revisa el escenario")' % (
            escenario_celda, fila_origen, fila_origen, fila_origen)
        f = fila_datos(ws, f, etiqueta, [formula], fmt=fmt)
        activos.append('Supuestos!$B$%d' % (f - 1))
    f += 1

    f = seccion(ws, f, 'Supuestos fijos')
    fijos = [('Depreciacion anual', 20, NUM_MILES),
             ('CapEx anual', 50, NUM_MILES),
             ('Ventas del Ano 0', 1000, NUM_MILES)]
    refs_fijos = []
    for etiqueta, valor, fmt in fijos:
        f = fila_datos(ws, f, etiqueta, [valor], fmt=fmt, input_=True)
        refs_fijos.append('Supuestos!$B$%d' % (f - 1))
    f += 1

    f = seccion(ws, f, 'Balance inicial (Ano 0)')
    balance0 = [('Efectivo', 100), ('Inventario', 200), ('PP&E neto', 800),
                ('Cuentas por pagar', 100), ('Deuda a largo plazo', 600),
                ('Ganancias retenidas', 400)]
    refs_b0 = []
    for etiqueta, valor in balance0:
        f = fila_datos(ws, f, etiqueta, [valor], fmt=NUM_MILES, input_=True)
        refs_b0.append('Supuestos!$B$%d' % (f - 1))

    return dict(crec=activos[0], margen=activos[1], impuestos=activos[2], interes=activos[3],
                deprec=refs_fijos[0], capex=refs_fijos[1], ventas0=refs_fijos[2],
                efectivo0=refs_b0[0], inventario0=refs_b0[1], ppe0=refs_b0[2],
                cxp0=refs_b0[3], deuda0=refs_b0[4], gr0=refs_b0[5])


def hoja_pl(wb, r):
    ws = wb.create_sheet('P&L')
    anchos(ws, 42, 14, 7)
    f = titulo(ws, 'ESTADO DE RESULTADOS PROYECTADO', 7)
    f = cabecera(ws, f, ['Concepto'] + ['Ano %d' % i for i in range(ANIOS + 1)])

    col = lambda i: chr(ord('B') + i)
    f = fila_datos(ws, f, 'Ventas',
                   ['=%s' % r['ventas0']] +
                   ['=%s%d*(1+%s)' % (col(i - 1), f, r['crec']) for i in range(1, ANIOS + 1)])
    fila_ventas = f - 1
    f = fila_datos(ws, f, 'EBITDA',
                   ['=%s%d*%s' % (col(i), fila_ventas, r['margen']) for i in range(ANIOS + 1)])
    fila_ebitda = f - 1
    f = fila_datos(ws, f, '(-) Depreciacion',
                   ['=-%s' % r['deprec']] * (ANIOS + 1))
    fila_dep = f - 1
    f = fila_datos(ws, f, 'EBIT',
                   ['=%s%d+%s%d' % (col(i), fila_ebitda, col(i), fila_dep) for i in range(ANIOS + 1)],
                   total=True)
    fila_ebit = f - 1
    f = fila_datos(ws, f, '(-) Gasto financiero',
                   ['=-%s*%s' % (r['deuda0'], r['interes'])] * (ANIOS + 1))
    fila_int = f - 1
    f = fila_datos(ws, f, 'EBT (utilidad antes de impuestos)',
                   ['=%s%d+%s%d' % (col(i), fila_ebit, col(i), fila_int) for i in range(ANIOS + 1)],
                   total=True)
    fila_ebt = f - 1
    f = fila_datos(ws, f, '(-) Impuestos',
                   ['=-MAX(0,%s%d)*%s' % (col(i), fila_ebt, r['impuestos']) for i in range(ANIOS + 1)])
    fila_imp = f - 1
    f = fila_datos(ws, f, 'UTILIDAD NETA',
                   ['=%s%d+%s%d' % (col(i), fila_ebt, col(i), fila_imp) for i in range(ANIOS + 1)],
                   total=True)
    fila_un = f - 1
    f += 1
    f = seccion(ws, f, 'Margenes de control')
    f = fila_datos(ws, f, 'Margen EBITDA',
                   ['=IFERROR(%s%d/%s%d,0)' % (col(i), fila_ebitda, col(i), fila_ventas)
                    for i in range(ANIOS + 1)], fmt=NUM_PCT)
    f = fila_datos(ws, f, 'Margen neto',
                   ['=IFERROR(%s%d/%s%d,0)' % (col(i), fila_un, col(i), fila_ventas)
                    for i in range(ANIOS + 1)], fmt=NUM_PCT)
    f += 1
    nota(ws, f, 'Los impuestos usan MAX(0;EBT): si hay perdida no se paga impuesto (no se modela el credito fiscal).')
    return dict(un=fila_un, ebit=fila_ebit, dep=fila_dep, ventas=fila_ventas)


def hoja_flujo(wb, r, pl):
    ws = wb.create_sheet('Flujo de Efectivo')
    anchos(ws, 42, 14, 7)
    f = titulo(ws, 'FLUJO DE EFECTIVO (metodo indirecto)', 7)
    f = cabecera(ws, f, ['Concepto'] + ['Ano %d' % i for i in range(ANIOS + 1)])
    col = lambda i: chr(ord('B') + i)

    f = fila_datos(ws, f, 'Utilidad neta',
                   [''] + ["='P&L'!%s%d" % (col(i), pl['un']) for i in range(1, ANIOS + 1)])
    fila_un = f - 1
    f = fila_datos(ws, f, '(+) Depreciacion (no consume caja)',
                   [''] + ['=%s' % r['deprec'] for _ in range(1, ANIOS + 1)])
    fila_dep = f - 1
    f = fila_datos(ws, f, 'Flujo de operacion (CFO)',
                   [''] + ['=%s%d+%s%d' % (col(i), fila_un, col(i), fila_dep)
                           for i in range(1, ANIOS + 1)], total=True)
    fila_cfo = f - 1
    f = fila_datos(ws, f, '(-) CapEx',
                   [''] + ['=-%s' % r['capex'] for _ in range(1, ANIOS + 1)])
    fila_capex = f - 1
    f = fila_datos(ws, f, 'Flujo de inversion (CFI)',
                   [''] + ['=%s%d' % (col(i), fila_capex) for i in range(1, ANIOS + 1)], total=True)
    fila_cfi = f - 1
    f = fila_datos(ws, f, 'Flujo de financiamiento (CFF)',
                   [''] + [0 for _ in range(1, ANIOS + 1)], total=True)
    fila_cff = f - 1
    f = nota(ws, f, 'Sin cambios en deuda ni dividendos por simplicidad: modelalos aqui cuando los necesites.')
    f = fila_datos(ws, f, 'VARIACION NETA DE EFECTIVO',
                   [''] + ['=%s%d+%s%d+%s%d' % (col(i), fila_cfo, col(i), fila_cfi, col(i), fila_cff)
                           for i in range(1, ANIOS + 1)], total=True)
    fila_var = f - 1
    f = fila_datos(ws, f, 'Efectivo inicial',
                   [''] + ['=%s' % r['efectivo0'] if i == 1 else '=%s%d' % (col(i - 1), f + 1)
                           for i in range(1, ANIOS + 1)])
    fila_ini = f - 1
    f = fila_datos(ws, f, 'EFECTIVO FINAL (el "plug" del balance)',
                   [''] + ['=%s%d+%s%d' % (col(i), fila_ini, col(i), fila_var)
                           for i in range(1, ANIOS + 1)], total=True)
    fila_fin = f - 1
    f += 1
    f = seccion(ws, f, 'Flujo de caja libre')
    f = fila_datos(ws, f, 'FCF = CFO - CapEx',
                   [''] + ['=%s%d+%s%d' % (col(i), fila_cfo, col(i), fila_capex)
                           for i in range(1, ANIOS + 1)], total=True)
    return dict(efectivo_final=fila_fin)


def hoja_balance(wb, r, pl, cf):
    ws = wb.create_sheet('Balance')
    anchos(ws, 42, 14, 7)
    f = titulo(ws, 'BALANCE GENERAL PROYECTADO', 7)
    f = cabecera(ws, f, ['Concepto'] + ['Ano %d' % i for i in range(ANIOS + 1)])
    col = lambda i: chr(ord('B') + i)

    f = seccion(ws, f, 'ACTIVO')
    f = fila_datos(ws, f, 'Efectivo',
                   ['=%s' % r['efectivo0']] +
                   ["='Flujo de Efectivo'!%s%d" % (col(i), cf['efectivo_final'])
                    for i in range(1, ANIOS + 1)])
    fila_ef = f - 1
    f = fila_datos(ws, f, 'Inventario', ['=%s' % r['inventario0']] * (ANIOS + 1))
    fila_inv = f - 1
    f = fila_datos(ws, f, 'PP&E neto',
                   ['=%s' % r['ppe0']] +
                   ['=%s%d+%s-%s' % (col(i - 1), f, r['capex'], r['deprec'])
                    for i in range(1, ANIOS + 1)])
    fila_ppe = f - 1
    f = fila_datos(ws, f, 'TOTAL ACTIVO',
                   ['=SUM(%s%d:%s%d)' % (col(i), fila_ef, col(i), fila_ppe)
                    for i in range(ANIOS + 1)], total=True)
    fila_ta = f - 1
    f += 1

    f = seccion(ws, f, 'PASIVO Y PATRIMONIO')
    f = fila_datos(ws, f, 'Cuentas por pagar', ['=%s' % r['cxp0']] * (ANIOS + 1))
    fila_cxp = f - 1
    f = fila_datos(ws, f, 'Deuda a largo plazo', ['=%s' % r['deuda0']] * (ANIOS + 1))
    fila_deuda = f - 1
    f = fila_datos(ws, f, 'Ganancias retenidas',
                   ['=%s' % r['gr0']] +
                   ["=%s%d+'P&L'!%s%d" % (col(i - 1), f, col(i), pl['un'])
                    for i in range(1, ANIOS + 1)])
    fila_gr = f - 1
    f = fila_datos(ws, f, 'TOTAL PASIVO + PATRIMONIO',
                   ['=SUM(%s%d:%s%d)' % (col(i), fila_cxp, col(i), fila_gr)
                    for i in range(ANIOS + 1)], total=True)
    fila_tp = f - 1
    f += 1

    f = seccion(ws, f, 'COMPROBACION')
    f = fila_datos(ws, f, 'Activo - (Pasivo + Patrimonio)  ->  debe ser 0',
                   ['=%s%d-%s%d' % (col(i), fila_ta, col(i), fila_tp)
                    for i in range(ANIOS + 1)], fmt=NUM_DEC, total=True)
    fila_chk = f - 1
    rango = 'B%d:%s%d' % (fila_chk, col(ANIOS), fila_chk)
    ws.conditional_formatting.add(rango, CellIsRule(
        operator='equal', formula=['0'],
        fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(color='006100', bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator='notEqual', formula=['0'],
        fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(color='9C0006', bold=True)))
    f += 1
    nota(ws, f, 'Verde = el balance cuadra. Rojo = hay un error de logica: NUNCA lo fuerces a mano.')


def main():
    wb = Workbook()
    wb.remove(wb.active)
    r = hoja_supuestos(wb)
    pl = hoja_pl(wb, r)
    cf = hoja_flujo(wb, r, pl)
    hoja_balance(wb, r, pl, cf)
    wb['Supuestos'].sheet_view.tabSelected = True
    wb.active = 0
    guardar(wb, 'modelo_3_estados.xlsx')


if __name__ == '__main__':
    main()
