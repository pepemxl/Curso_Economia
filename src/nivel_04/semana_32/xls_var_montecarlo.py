# -*- coding: utf-8 -*-
"""Semanas 20 y 32: VaR parametrico, cobertura con futuros y simulacion Monte Carlo.

Calibrado con los ejercicios del curso: VaR 1d = 232.600, VaR 10d = 735.546,
24 contratos de cobertura, y el proyecto inmobiliario con ~55% de exito.
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
from openpyxl import Workbook
from openpyxl.styles import Font
from estilo_excel import (titulo, seccion, cabecera, fila_datos, nota, anchos, guardar,
                          NUM_MILES, NUM_DEC, NUM_PCT, NUM_PCT2, BORDE, R_INPUT)

ITERACIONES = 1000


def hoja_var(wb):
    ws = wb.create_sheet('VaR parametrico')
    anchos(ws, 52, 16, 3)
    f = titulo(ws, 'VALUE AT RISK PARAMETRICO', 3)

    f = seccion(ws, f, 'Datos del portafolio (amarillo = editable)')
    f = fila_datos(ws, f, 'Valor del portafolio', [5000000], fmt=NUM_MILES, input_=True)
    r_v = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Volatilidad diaria (sigma)', [0.02], fmt=NUM_PCT2, input_=True)
    r_s = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Nivel de confianza', [0.99], fmt=NUM_PCT, input_=True)
    r_c = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Z correspondiente  =NORM.S.INV(confianza)',
                   ['=NORM.S.INV(%s)' % r_c], fmt=NUM_DEC)
    r_z = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Horizonte en dias', [10], fmt='0', input_=True)
    r_t = 'B%d' % (f - 1)
    f += 1

    f = seccion(ws, f, 'Resultados')
    f = fila_datos(ws, f, 'Desviacion estandar en dinero  V x sigma',
                   ['=%s*%s' % (r_v, r_s)], fmt=NUM_MILES)
    r_sd = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'VaR a 1 dia  =  V x Z x sigma',
                   ['=%s*%s' % (r_sd, r_z)], fmt=NUM_MILES, total=True)
    r_var1 = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'VaR a T dias  =  VaR_1d x RAIZ(T)',
                   ['=%s*SQRT(%s)' % (r_var1, r_t)], fmt=NUM_MILES, total=True)
    f = fila_datos(ws, f, 'VaR a 1 dia como % del portafolio',
                   ['=%s/%s' % (r_var1, r_v)], fmt=NUM_PCT2)
    f += 1
    f = fila_datos(ws, f, 'Expected Shortfall (CVaR) al mismo nivel',
                   ['=%s*NORM.DIST(%s,0,1,FALSE)/(1-%s)' % (r_sd, r_z, r_c)],
                   fmt=NUM_MILES, total=True)
    f = nota(ws, f, 'El ES responde "si estoy en el peor 1%, cuanto pierdo en promedio": siempre mayor que el VaR.')
    f = nota(ws, f, 'Se escala con RAIZ(T), no con T: la varianza crece con el tiempo, la desviacion con su raiz.')
    f += 1

    f = seccion(ws, f, 'Cobertura con futuros del indice')
    f = fila_datos(ws, f, 'Beta del portafolio respecto al indice', [1.2], fmt=NUM_DEC, input_=True)
    r_b = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Valor nocional de un contrato de futuro', [250000], fmt=NUM_MILES, input_=True)
    r_cf = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Exposicion ajustada por beta  V x beta',
                   ['=%s*%s' % (r_v, r_b)], fmt=NUM_MILES)
    r_exp = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'CONTRATOS A VENDER (posicion corta)',
                   ['=IFERROR(%s/%s,"Revisa el nocional")' % (r_exp, r_cf)],
                   fmt=NUM_DEC, total=True)
    ws['B%d' % (f - 1)].font = Font(bold=True, size=13)
    f = nota(ws, f, 'Los contratos son indivisibles: si sale 24,7 tendras que elegir 24 o 25 y aceptar cobertura imperfecta.')
    return ws


def hoja_montecarlo(wb):
    ws = wb.create_sheet('Monte Carlo')
    anchos(ws, 44, 16, 6)
    f = titulo(ws, 'SIMULACION MONTE CARLO  ·  proyecto inmobiliario', 6)

    f = seccion(ws, f, 'Supuestos (amarillo = editable)')
    f = fila_datos(ws, f, 'Inversion inicial (Ano 0)', [-2000000], fmt=NUM_MILES, input_=True)
    r_inv = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Precio de venta esperado (media)', [2500000], fmt=NUM_MILES, input_=True)
    r_mv = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Desviacion del precio de venta', [300000], fmt=NUM_MILES, input_=True)
    r_sv = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Costo de construccion esperado (media)', [300000], fmt=NUM_MILES, input_=True)
    r_mc = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Desviacion del costo', [50000], fmt=NUM_MILES, input_=True)
    r_sc = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Tasa de descuento', [0.08], fmt=NUM_PCT2, input_=True)
    r_tasa = 'B%d' % (f - 1)
    f += 1

    f = seccion(ws, f, 'Resultado teorico (para contrastar con la simulacion)')
    f = fila_datos(ws, f, 'VAN esperado  =  I0 + (media venta - media costo)/(1+r)',
                   ['=%s+(%s-%s)/(1+%s)' % (r_inv, r_mv, r_mc, r_tasa)], fmt=NUM_DEC)
    f = fila_datos(ws, f, 'Sigma del flujo  =  RAIZ(sv^2 + sc^2)',
                   ['=SQRT(%s^2+%s^2)' % (r_sv, r_sc)], fmt=NUM_MILES)
    r_sf = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'P(VAN > 0) teorica',
                   ['=1-NORM.DIST(-%s*(1+%s),%s-%s,%s,TRUE)' % (r_inv, r_tasa, r_mv, r_mc, r_sf)],
                   fmt=NUM_PCT2)
    f = nota(ws, f, 'Al restar dos normales independientes las VARIANZAS se suman, no las desviaciones.')
    f += 1

    f = seccion(ws, f, 'Estadisticos de las %d iteraciones' % ITERACIONES)
    fila_est = f
    prim = fila_est + 12          # primera fila de la simulacion
    ult = prim + ITERACIONES - 1
    rng = '$E$%d:$E$%d' % (prim, ult)
    f = fila_datos(ws, f, 'VAN medio simulado', ['=AVERAGE(%s)' % rng], fmt=NUM_DEC, total=True)
    f = fila_datos(ws, f, 'Desviacion estandar del VAN', ['=STDEV.S(%s)' % rng], fmt=NUM_DEC)
    f = fila_datos(ws, f, 'VAN minimo', ['=MIN(%s)' % rng], fmt=NUM_DEC)
    f = fila_datos(ws, f, 'VAN maximo', ['=MAX(%s)' % rng], fmt=NUM_DEC)
    f = fila_datos(ws, f, 'Iteraciones con VAN > 0', ['=COUNTIF(%s,">0")' % rng], fmt=NUM_MILES)
    f = fila_datos(ws, f, 'PROBABILIDAD DE EXITO',
                   ['=COUNTIF(%s,">0")/%d' % (rng, ITERACIONES)], fmt=NUM_PCT2, total=True)
    ws['B%d' % (f - 1)].font = Font(bold=True, size=13)
    f = fila_datos(ws, f, 'VaR del proyecto al 95% (percentil 5)',
                   ['=PERCENTILE.INC(%s,0.05)' % rng], fmt=NUM_DEC)
    f += 1
    f = nota(ws, f, 'Pulsa F9 para volver a sortear: ALEATORIO() es volatil y cada recalculo genera 1.000 escenarios nuevos.')

    f = cabecera(ws, prim - 1, ['Iteracion', 'Precio de venta', 'Costo', 'Flujo Ano 1', 'VAN'])
    for i in range(ITERACIONES):
        r = prim + i
        ws.cell(r, 1, i + 1).number_format = '0'
        ws.cell(r, 2, '=NORM.INV(RAND(),$%s,$%s)' % (r_mv, r_sv)).number_format = NUM_MILES
        ws.cell(r, 3, '=NORM.INV(RAND(),$%s,$%s)' % (r_mc, r_sc)).number_format = NUM_MILES
        ws.cell(r, 4, '=B%d-C%d' % (r, r)).number_format = NUM_MILES
        ws.cell(r, 5, '=$%s+D%d/(1+$%s)' % (r_inv, r, r_tasa)).number_format = NUM_MILES
    ws.freeze_panes = 'A%d' % prim
    return dict(prim=prim, ult=ult)


def main():
    wb = Workbook(); wb.remove(wb.active)
    hoja_var(wb)
    hoja_montecarlo(wb)
    wb.active = 0
    guardar(wb, 'var_montecarlo.xlsx')


if __name__ == '__main__':
    main()
