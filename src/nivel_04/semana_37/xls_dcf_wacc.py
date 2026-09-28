# -*- coding: utf-8 -*-
"""Semanas 24 y 37: calculo del WACC y valuacion por DCF con sensibilidad.

Calibrado con los ejercicios del curso: los datos de GoldRush (Semana 24) dan
WACC = 8,00% y los de LogiTrans (Semana 37) dan $116,79 por accion.
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from estilo_excel import (titulo, seccion, cabecera, fila_datos, nota, anchos, guardar,
                          NUM_MILES, NUM_DEC, NUM_PCT, NUM_PCT2, R_INPUT, R_TOTAL,
                          F_SECCION, F_TOTAL, BORDE)

ANIOS = 3


def hoja_wacc(wb):
    ws = wb.create_sheet('WACC')
    anchos(ws, 46, 16, 4)
    f = titulo(ws, 'COSTO PROMEDIO PONDERADO DE CAPITAL (WACC)', 4)

    f = seccion(ws, f, 'Datos de entrada (amarillo = editable)')
    entradas = [('Tasa libre de riesgo (Rf)', 0.035, NUM_PCT2),
                ('Beta apalancada de la empresa', 1.1, NUM_DEC),
                ('Prima por riesgo de mercado (MRP)', 0.055, NUM_PCT2),
                ('Costo de la deuda antes de impuestos (rd)', 0.07, NUM_PCT2),
                ('Tasa de impuestos', 0.30, NUM_PCT),
                ('Valor de mercado del patrimonio (E)', 400, NUM_MILES),
                ('Valor de mercado de la deuda (D)', 200, NUM_MILES)]
    ref = {}
    claves = ['rf', 'beta', 'mrp', 'rd', 'tax', 'E', 'D']
    for (etiqueta, valor, fmt), clave in zip(entradas, claves):
        f = fila_datos(ws, f, etiqueta, [valor], fmt=fmt, input_=True)
        ref[clave] = 'B%d' % (f - 1)
    f += 1

    f = seccion(ws, f, 'Calculos')
    f = fila_datos(ws, f, 'Costo del patrimonio  re = Rf + beta x MRP  (CAPM)',
                   ['=%s+%s*%s' % (ref['rf'], ref['beta'], ref['mrp'])], fmt=NUM_PCT2)
    ref['re'] = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Costo de la deuda despues de impuestos  rd x (1-t)',
                   ['=%s*(1-%s)' % (ref['rd'], ref['tax'])], fmt=NUM_PCT2)
    ref['rd_neto'] = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Valor total de la empresa  V = E + D',
                   ['=%s+%s' % (ref['E'], ref['D'])], fmt=NUM_MILES)
    ref['V'] = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Peso del patrimonio  E/V',
                   ['=%s/%s' % (ref['E'], ref['V'])], fmt=NUM_PCT)
    ref['wE'] = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Peso de la deuda  D/V',
                   ['=%s/%s' % (ref['D'], ref['V'])], fmt=NUM_PCT)
    ref['wD'] = 'B%d' % (f - 1)
    f += 1
    f = fila_datos(ws, f, 'WACC',
                   ['=%s*%s+%s*%s' % (ref['wE'], ref['re'], ref['wD'], ref['rd_neto'])],
                   fmt=NUM_PCT2, total=True)
    ref['wacc'] = 'B%d' % (f - 1)
    ws['B%d' % (f - 1)].font = Font(bold=True, size=14)
    f += 1
    f = nota(ws, f, 'Los pesos deben ir a valor de MERCADO, nunca contable.')
    f = nota(ws, f, 'Si subes la deuda, el beta del patrimonio deberia subir tambien: el WACC no baja indefinidamente.')
    return ref


def hoja_dcf(wb, wacc_ref):
    ws = wb.create_sheet('DCF')
    anchos(ws, 46, 16, 6)
    f = titulo(ws, 'VALUACION POR DESCUENTO DE FLUJOS (DCF)', 6)

    f = seccion(ws, f, 'Supuestos')
    # El WACC puede venir de la hoja WACC (misma empresa) o introducirse a mano
    # (por ejemplo al valorar una empresa distinta de la calculada alli).
    f = fila_datos(ws, f, 'Origen del WACC  (1 = hoja WACC   2 = manual)', [2],
                   fmt='0', input_=True)
    r_origen = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'WACC calculado en la hoja WACC', ['=WACC!%s' % wacc_ref],
                   fmt=NUM_PCT2)
    r_wacc_hoja = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'WACC manual', [0.12], fmt=NUM_PCT2, input_=True)
    r_wacc_man = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'WACC EN USO',
                   ['=IFERROR(CHOOSE(%s,%s,%s),"Origen debe ser 1 o 2")'
                    % (r_origen, r_wacc_hoja, r_wacc_man)], fmt=NUM_PCT2, total=True)
    r_wacc = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Crecimiento perpetuo g', [0.015], fmt=NUM_PCT2, input_=True)
    r_g = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Deuda neta', [40], fmt=NUM_MILES, input_=True)
    r_dn = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Acciones en circulacion (millones)', [5], fmt=NUM_DEC, input_=True)
    r_acc = 'B%d' % (f - 1)
    f += 1

    f = seccion(ws, f, 'Flujos explicitos')
    f = cabecera(ws, f, ['Concepto'] + ['Ano %d' % i for i in range(1, ANIOS + 1)])
    f = fila_datos(ws, f, 'FCFF proyectado', [50, 60, 70], fmt=NUM_MILES, input_=True)
    r_fcff = f - 1
    col = lambda i: chr(ord('B') + i)
    f = fila_datos(ws, f, 'Factor de descuento  1/(1+WACC)^t',
                   ['=1/(1+$%s)^%d' % (r_wacc, i + 1) for i in range(ANIOS)], fmt='0.0000')
    r_fd = f - 1
    f = fila_datos(ws, f, 'Valor presente del FCFF',
                   ['=%s%d*%s%d' % (col(i), r_fcff, col(i), r_fd) for i in range(ANIOS)],
                   fmt=NUM_DEC)
    r_vp = f - 1
    f = fila_datos(ws, f, 'Suma de VP explicitos',
                   ['=SUM(B%d:%s%d)' % (r_vp, col(ANIOS - 1), r_vp)], fmt=NUM_DEC, total=True)
    r_suma = 'B%d' % (f - 1)
    f += 1

    f = seccion(ws, f, 'Valor terminal (Gordon)')
    f = fila_datos(ws, f, 'FCFF del ano %d = FCFF%d x (1+g)' % (ANIOS + 1, ANIOS),
                   ['=%s%d*(1+$%s)' % (col(ANIOS - 1), r_fcff, r_g)], fmt=NUM_DEC)
    r_fcff4 = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'Valor terminal al final del ano %d' % ANIOS,
                   ['=IFERROR(%s/($%s-$%s),"g debe ser menor que el WACC")' % (r_fcff4, r_wacc, r_g)],
                   fmt=NUM_DEC)
    r_tv = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'VP del valor terminal (descontado %d anos, NO %d)' % (ANIOS, ANIOS + 1),
                   ['=%s*%s%d' % (r_tv, col(ANIOS - 1), r_fd)], fmt=NUM_DEC, total=True)
    r_vptv = 'B%d' % (f - 1)
    f = nota(ws, f, 'Error clasico: descontar el VT un periodo de mas. Gordon lo situa en el ano n, no en n+1.')
    f += 1

    f = seccion(ws, f, 'Puente de valuacion')
    f = fila_datos(ws, f, 'Enterprise Value', ['=%s+%s' % (r_suma, r_vptv)],
                   fmt=NUM_DEC, total=True)
    r_ev = 'B%d' % (f - 1)
    f = fila_datos(ws, f, '(-) Deuda neta', ['=-%s' % r_dn], fmt=NUM_DEC)
    f = fila_datos(ws, f, 'Equity Value', ['=%s-%s' % (r_ev, r_dn)], fmt=NUM_DEC, total=True)
    r_eq = 'B%d' % (f - 1)
    f = fila_datos(ws, f, 'VALOR INTRINSECO POR ACCION',
                   ['=IFERROR(%s/%s,"Revisa el numero de acciones")' % (r_eq, r_acc)],
                   fmt=NUM_DEC, total=True)
    r_precio = 'B%d' % (f - 1)
    ws['B%d' % (f - 1)].font = Font(bold=True, size=14)
    f += 1
    f = fila_datos(ws, f, 'Peso del valor terminal sobre el EV',
                   ['=%s/%s' % (r_vptv, r_ev)], fmt=NUM_PCT)
    f = nota(ws, f, 'Si el VT supera el 80% del EV, el modelo depende mas de un supuesto que del analisis.')
    return dict(wacc=r_wacc, g=r_g, precio=r_precio, ev=r_ev,
                fila_fcff=r_fcff, dn=r_dn, acc=r_acc, suma=r_suma)


def hoja_sensibilidad(wb, d):
    """Matriz WACC x g construida con formulas explicitas (auditable, sin tabla de datos)."""
    ws = wb.create_sheet('Sensibilidad')
    anchos(ws, 22, 13, 8)
    f = titulo(ws, 'SENSIBILIDAD DEL PRECIO POR ACCION:  WACC  x  crecimiento perpetuo (g)', 8)
    f = nota(ws, f, 'Cada celda recalcula el DCF completo con ese par (WACC, g). Un DCF serio se presenta como rango, no como un numero.')
    f += 1

    gs = [0.005, 0.010, 0.015, 0.020, 0.025, 0.030]
    waccs = [0.10, 0.11, 0.12, 0.13, 0.14, 0.15]

    fila_cab = f
    ws.cell(fila_cab, 1, 'WACC \\ g').font = F_SECCION
    for j, g in enumerate(gs):
        c = ws.cell(fila_cab, 2 + j, g)
        c.number_format = NUM_PCT2
        c.font = F_TOTAL
        c.alignment = Alignment(horizontal='center')
        c.fill = R_TOTAL
    f += 1

    prim = f
    for i, w in enumerate(waccs):
        c = ws.cell(f, 1, w)
        c.number_format = NUM_PCT2
        c.font = F_TOTAL
        c.fill = R_TOTAL
        for j, g in enumerate(gs):
            wc = '$A%d' % f          # WACC de la fila
            gc = '%s$%d' % (chr(ord('B') + j), fila_cab)   # g de la columna
            # VP explicitos + VP del valor terminal, todo en funcion de (WACC, g)
            vp = '+'.join('DCF!%s%d/(1+%s)^%d' % (chr(ord('B') + t), d['fila_fcff'], wc, t + 1)
                          for t in range(ANIOS))
            tv = '(DCF!%s%d*(1+%s)/(%s-%s))/(1+%s)^%d' % (
                chr(ord('B') + ANIOS - 1), d['fila_fcff'], gc, wc, gc, wc, ANIOS)
            formula = '=IFERROR((%s+%s-DCF!$%s)/DCF!$%s,"n/a")' % (vp, tv, d['dn'], d['acc'])
            cel = ws.cell(f, 2 + j, formula)
            cel.number_format = NUM_DEC
            cel.border = BORDE
        f += 1

    rango = 'B%d:%s%d' % (prim, chr(ord('B') + len(gs) - 1), f - 1)
    ws.conditional_formatting.add(rango, ColorScaleRule(
        start_type='min', start_color='F8696B',
        mid_type='percentile', mid_value=50, mid_color='FFEB84',
        end_type='max', end_color='63BE7B'))
    return dict(primera=prim, fila_cab=fila_cab)


def main():
    wb = Workbook(); wb.remove(wb.active)
    ref = hoja_wacc(wb)
    d = hoja_dcf(wb, ref['wacc'])
    hoja_sensibilidad(wb, d)
    wb.active = 0
    guardar(wb, 'dcf_wacc.xlsx')


if __name__ == '__main__':
    main()
