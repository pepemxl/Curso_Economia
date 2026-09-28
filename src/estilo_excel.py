# -*- coding: utf-8 -*-
"""Estilo y utilidades comunes para las plantillas de Excel del curso.

Se generan con `make plantillas`, que ejecuta cada `src/**/xls_*.py` y deja
los .xlsx en `docs/assets/plantillas/`.

IMPORTANTE sobre los nombres de función
---------------------------------------
El formato .xlsx guarda las fórmulas **siempre en inglés** (NPV, IRR, PMT,
CHOOSE...). Excel las muestra traducidas segun el idioma de la instalacion
(VNA, TIR, PAGO, ELEGIR...). Por eso los scripts escriben nombres en ingles:
al abrir el archivo en un Excel en español apareceran ya en español.
"""
import os
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta alineada con la de las figuras
AZUL      = '1F77B4'
ROJO      = 'D62728'
VERDE     = '2CA02C'
NARANJA   = 'FF7F0E'
GRIS      = 'F2F2F2'
GRIS_OSC  = '7F7F7F'
AMARILLO  = 'FFF2CC'   # convencion de banca: celda de input editable

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR_PLANTILLAS = os.path.join(RAIZ, 'docs', 'assets', 'plantillas')

F_TITULO   = Font(bold=True, size=14, color='FFFFFF')
F_CABECERA = Font(bold=True, size=11, color='FFFFFF')
F_SECCION  = Font(bold=True, size=11, color=AZUL)
F_NORMAL   = Font(size=11)
F_TOTAL    = Font(bold=True, size=11)
F_NOTA     = Font(size=9, italic=True, color=GRIS_OSC)

R_TITULO   = PatternFill('solid', fgColor=AZUL)
R_CABECERA = PatternFill('solid', fgColor=GRIS_OSC)
R_INPUT    = PatternFill('solid', fgColor=AMARILLO)
R_TOTAL    = PatternFill('solid', fgColor=GRIS)

_fino = Side(style='thin', color='BFBFBF')
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
BORDE_SUP = Border(top=Side(style='medium', color='404040'))

NUM_MILES   = '#,##0'
NUM_DEC     = '#,##0.00'
NUM_PCT     = '0.0%'
NUM_PCT2    = '0.00%'
NUM_DINERO  = '#,##0.00 $'


def titulo(ws, texto, ancho=6, fila=1):
    """Barra de titulo que ocupa las primeras `ancho` columnas."""
    ws.cell(fila, 1, texto).font = F_TITULO
    for c in range(1, ancho + 1):
        ws.cell(fila, c).fill = R_TITULO
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    ws.row_dimensions[fila].height = 24
    return fila + 2


def seccion(ws, fila, texto):
    ws.cell(fila, 1, texto).font = F_SECCION
    return fila + 1


def cabecera(ws, fila, valores, col_inicio=1):
    for i, v in enumerate(valores):
        c = ws.cell(fila, col_inicio + i, v)
        c.font, c.fill, c.border = F_CABECERA, R_CABECERA, BORDE
        c.alignment = Alignment(horizontal='center')
    return fila + 1


def fila_datos(ws, fila, etiqueta, valores, fmt=NUM_MILES,
               total=False, input_=False, col_inicio=2):
    """Escribe una fila: etiqueta + valores (numeros o formulas '=...')."""
    c = ws.cell(fila, 1, etiqueta)
    c.font = F_TOTAL if total else F_NORMAL
    if total:
        c.fill = R_TOTAL
    for i, v in enumerate(valores):
        cel = ws.cell(fila, col_inicio + i, v)
        cel.number_format = fmt
        cel.border = BORDE
        cel.font = F_TOTAL if total else F_NORMAL
        if total:
            cel.fill = R_TOTAL
        if input_:
            cel.fill = R_INPUT
    return fila + 1


def nota(ws, fila, texto, col=1):
    ws.cell(fila, col, texto).font = F_NOTA
    return fila + 1


def anchos(ws, primera=42, resto=14, n=10):
    ws.column_dimensions['A'].width = primera
    for i in range(2, n + 2):
        ws.column_dimensions[get_column_letter(i)].width = resto


def guardar(wb, nombre):
    if not os.path.isdir(DIR_PLANTILLAS):
        os.makedirs(DIR_PLANTILLAS)
    destino = os.path.join(DIR_PLANTILLAS, nombre)
    wb.save(destino)
    print('  -> %s' % os.path.relpath(destino, RAIZ))
